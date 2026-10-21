"""
Forwarding in the portal (Documentation/adr/0018-forwarding.md): the resolver of last resort (off with GS1's address
until set; switched on and off; its address checked — https only, never this installation or one of its verified
domains), the prefix rules (added, changed, switched off, removed; every refusal; the organisations whose prefixes
they overlap), who may change them (server administrators outside support access), the journal, the settings as the
resolver's own reader (web_server/src/forwarding.py) sees them, the request log's two outcomes, and the overview's
figures: the quality counts for every scope, the Forwarding group of server administrators with the requests by rule,
and the exports.

  pip install -r portal/requirements.txt mongomock
  python dev-tests/portal/test_forwarding.py
"""
import csv
import importlib.util
import io
import os
import sys
import tempfile
import threading
from datetime import datetime, timedelta, timezone

from werkzeug.serving import make_server

HERE = os.path.dirname(os.path.abspath(__file__))
PORTAL = os.environ.get("PORTAL_DIR", os.path.join(HERE, "..", "..", "portal"))
CONFIG = tempfile.mkdtemp()
DATA_ENTRY_PORT = 3419
sys.path[:0] = [HERE, PORTAL]

os.environ.update(DATA_ENTRY_URL=f"http://127.0.0.1:{DATA_ENTRY_PORT}/api", SESSION_TOKEN="tok",
                  PORTAL_CONFIG_DIR=CONFIG, PORTAL_ORIGIN="https://id.example.org", PORTAL_COOKIE_SECURE="false",
                  RESOLVER_PUBLIC_URL="https://id.example.org", PORTAL_TIME_ZONE="America/Sao_Paulo")
os.environ["PORTAL_DB_URI"] = "mongomock://"
os.environ["PORTAL_DB_NAME"] = f"devtest_forwarding_{os.getpid()}"
os.environ["PORTAL_WORKER"] = "external"
os.environ["PORTAL_MONITOR_MINUTES"] = "0"
for name in ("RESOLVER_LOG_DAYS", "PORTAL_MAIL_HOST"):
    os.environ.pop(name, None)

import mock_data_entry  # noqa: E402
import store  # noqa: E402
import users  # noqa: E402
import tenants  # noqa: E402
users.create("root", "root-password-long", "admin", must_change=False)
import app as portal  # noqa: E402
import forwarding  # noqa: E402
import requestlog  # noqa: E402

threading.Thread(target=make_server("127.0.0.1", DATA_ENTRY_PORT, mock_data_entry.app, threaded=True).serve_forever,
                 daemon=True).start()
failures = []
ORIGIN = {"Origin": "https://id.example.org"}


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + ("" if condition else f"  → {detail}"))
    if not condition:
        failures.append(name)


def client(username=None, password=None):
    c = portal.app.test_client()
    if username:
        c.post("/portal/api/login", json={"username": username, "password": password}, headers=ORIGIN)
    return c


def code(r):
    return (r.get_json(silent=True) or {}).get("code")


acme = tenants.create("Acme Foods", ["7891234"], "root")
beta = tenants.create("Beta Drinks", ["7895555"], "root")
users.create("ana", "ana-password-long", "admin", tenant=acme["id"], must_change=False)
store.col("domains").insert_one({"_id": "id.acme.example", "tenant": acme["id"], "verified": True, "mode": "direct"})
root, ana, anyone = client("root", "root-password-long"), client("ana", "ana-password-long"), client()

# The resolver's own reader of the same settings (web_server/src/forwarding.py), on the portal's database
RESOLVER_SRC = os.path.join(HERE, "..", "..", "web_server", "src")
sys.path.insert(0, RESOLVER_SRC)
spec = importlib.util.spec_from_file_location("resolver_forwarding", os.path.join(RESOLVER_SRC, "forwarding.py"))
resolver_forwarding = importlib.util.module_from_spec(spec)
spec.loader.exec_module(resolver_forwarding)
sys.path.remove(RESOLVER_SRC)
reader = resolver_forwarding.Forwarding(lambda: store.db())


def seen():
    reader.forget()
    return reader


# ================================================================== the resolver of last resort
data = root.get("/portal/api/forwarding").get_json()
check("nothing set: the resolver of last resort off, GS1's address offered, no rule", data["fallback"]["enabled"] is False
      and data["fallback"]["url"] == "https://id.gs1.org" and data["fallback"]["default"] == "https://id.gs1.org"
      and data["rules"] == [] and data["readOnly"] is False, data)
check("…the resolver sees it off", seen().fallback() is None)
r = root.put("/portal/api/forwarding/fallback", json={"enabled": True, "url": ""}, headers=ORIGIN)
check("switched on without an address: GS1's global resolver", r.status_code == 200
      and code(r) == "forwarding.fallbackSaved" and r.get_json()["fallback"]["url"] == "https://id.gs1.org"
      and r.get_json()["fallback"]["enabled"] is True and r.get_json()["fallback"]["updatedBy"] == "root", r.get_json())
check("…the resolver sends unknown keys there", seen().fallback() == "https://id.gs1.org")
entry = store.col("journal").find_one({"action": "forwarding-fallback"}) or {}
check("…in the journal", entry.get("user") == "root" and entry.get("detail") == "enabled=true url=https://id.gs1.org", entry)
r = root.put("/portal/api/forwarding/fallback", json={"enabled": True, "url": " https://resolver.example.net/gs1 "},
             headers=ORIGIN)
check("…another address, trimmed", r.get_json()["fallback"]["url"] == "https://resolver.example.net/gs1"
      and seen().fallback() == "https://resolver.example.net/gs1", r.get_json())
for name, body, expected in (
        ("http", {"enabled": True, "url": "http://resolver.example.net"}, "forwarding.urlInvalid"),
        ("no host", {"enabled": True, "url": "https:///01/1"}, "forwarding.urlInvalid"),
        ("a password in it", {"enabled": True, "url": "https://u:p@resolver.example.net"}, "forwarding.urlInvalid"),
        ("spaces", {"enabled": True, "url": "https://resolver.example.net/a b"}, "forwarding.urlInvalid"),
        ("too long", {"enabled": True, "url": "https://r.example.net/" + "a" * 500}, "forwarding.urlTooLong"),
        ("this installation", {"enabled": True, "url": "https://ID.example.org/"}, "forwarding.urlSelf"),
        ("an organisation's verified domain", {"enabled": True, "url": "https://id.acme.example"}, "forwarding.urlSelf"),
        ("on as text", {"enabled": "yes", "url": ""}, "forwarding.enabledInvalid")):
    r = root.put("/portal/api/forwarding/fallback", json=body, headers=ORIGIN)
    check(f"…refused: {name} → {expected}", r.status_code == 422 and code(r) == expected, (r.status_code, r.get_json()))
check("…a refusal changes nothing", forwarding.fallback()["url"] == "https://resolver.example.net/gs1")
check("…the loop named", root.put("/portal/api/forwarding/fallback", json={"enabled": True, "url": "https://id.example.org"},
                                  headers=ORIGIN).get_json()["params"] == {"host": "id.example.org"})
r = root.put("/portal/api/forwarding/fallback", json={"enabled": False, "url": "https://resolver.example.net/gs1"},
             headers=ORIGIN)
check("switched off: the address kept for next time", code(r) == "forwarding.fallbackOff"
      and r.get_json()["fallback"]["url"] == "https://resolver.example.net/gs1" and seen().fallback() is None)

# ================================================================== who may
check("not without signing in", anyone.get("/portal/api/forwarding").status_code == 401)
r = ana.get("/portal/api/forwarding")
check("not an organisation's administrator (server administrators only)", r.status_code == 403, r.status_code)
check("…nor change it", ana.put("/portal/api/forwarding/fallback", json={"enabled": True}, headers=ORIGIN).status_code
      == 403 and ana.post("/portal/api/forwarding/rules", json={"prefix": "952"}, headers=ORIGIN).status_code == 403)
root.post("/portal/api/support", json={"tenant": acme["id"], "reason": "Ticket 4"}, headers=ORIGIN)
r = root.put("/portal/api/forwarding/fallback", json={"enabled": True}, headers=ORIGIN)
check("support access: read only (403 forwarding.readOnly)", r.status_code == 403 and code(r) == "forwarding.readOnly"
      and root.get("/portal/api/forwarding").get_json()["readOnly"] is True, (r.status_code, r.get_json()))
check("…no rule either", code(root.post("/portal/api/forwarding/rules", json={"prefix": "952", "target": "https://x.example"},
                                        headers=ORIGIN)) == "forwarding.readOnly")
root.delete("/portal/api/support", headers=ORIGIN)

# ================================================================== prefix rules
r = root.post("/portal/api/forwarding/rules", json={"prefix": "9521234", "target": "https://other.example.net",
                                                    "note": "  Moved   to another resolver "}, headers=ORIGIN)
rule = next((x for x in (r.get_json() or {}).get("rules", []) if x["prefix"] == "9521234"), {})
check("a rule: the path kept and on by default, the note tidied", r.status_code == 200 and code(r) == "forwarding.ruleSaved"
      and (rule.get("target"), rule.get("mode"), rule.get("enabled"), rule.get("note"), rule.get("organisations"))
      == ("https://other.example.net", "path", True, "Moved to another resolver", []), r.get_json())
check("…the resolver forwards that prefix", seen().rule_for("01", "09521234000019") == ("9521234", "https://other.example.net", "path")
      and seen().rule_for("01", "07891234500014") is None)
entry = store.col("journal").find_one({"action": "forwarding-rule-create"}) or {}
check("…in the journal", entry.get("detail") == "prefix=9521234 target=https://other.example.net enabled=true", entry)
for name, body, expected in (
        ("the same prefix again", {"prefix": "9521234", "target": "https://x.example"}, "forwarding.prefixExists"),
        ("two digits", {"prefix": "95", "target": "https://x.example"}, "forwarding.prefixInvalid"),
        ("thirteen digits", {"prefix": "9521234000019", "target": "https://x.example"}, "forwarding.prefixInvalid"),
        ("letters", {"prefix": "95A", "target": "https://x.example"}, "forwarding.prefixInvalid"),
        ("no destination", {"prefix": "952"}, "forwarding.urlRequired"),
        ("an http destination", {"prefix": "952", "target": "http://x.example"}, "forwarding.urlInvalid"),
        ("this installation", {"prefix": "952", "target": "https://id.example.org/x"}, "forwarding.urlSelf"),
        ("another mode", {"prefix": "952", "target": "https://x.example", "mode": "proxy"}, "forwarding.modeInvalid"),
        ("on as text", {"prefix": "952", "target": "https://x.example", "enabled": "on"}, "forwarding.enabledInvalid"),
        ("a note too long", {"prefix": "952", "target": "https://x.example", "note": "n" * 201}, "forwarding.noteTooLong")):
    r = root.post("/portal/api/forwarding/rules", json=body, headers=ORIGIN)
    check(f"…refused: {name} → {expected}", r.status_code == 422 and code(r) == expected, (r.status_code, r.get_json()))
check("…nothing stored by a refusal", store.col("forwarding_rules").count_documents({}) == 1)
r = root.post("/portal/api/forwarding/rules", json={"prefix": "789", "target": "https://brazil.example.net/scan",
                                                    "mode": "fixed"}, headers=ORIGIN)
rule = next((x for x in r.get_json()["rules"] if x["prefix"] == "789"), {})
check("a rule over organisations' prefixes: saved, the organisations named", rule.get("organisations")
      == ["Acme Foods", "Beta Drinks"] and rule.get("mode") == "fixed", rule)
check("…the list by prefix", [x["prefix"] for x in r.get_json()["rules"]] == ["789", "9521234"])
r = root.put("/portal/api/forwarding/rules/789", json={"target": "https://brazil.example.net/scan", "mode": "fixed",
                                                       "enabled": False}, headers=ORIGIN)
check("switched off: kept, the resolver ignores it", r.status_code == 200
      and next(x for x in r.get_json()["rules"] if x["prefix"] == "789")["enabled"] is False
      and seen().rule_for("01", "07891234500014") is None, r.get_json())
entry = store.col("journal").find_one({"action": "forwarding-rule-update"}) or {}
check("…in the journal", entry.get("detail") == "prefix=789 target=https://brazil.example.net/scan enabled=false", entry)
r = root.put("/portal/api/forwarding/rules/7895555", json={"target": "https://x.example"}, headers=ORIGIN)
check("changing a rule that does not exist: 404 forwarding.unknownRule", r.status_code == 404
      and code(r) == "forwarding.unknownRule" and r.get_json()["params"] == {"prefix": "7895555"}, r.get_json())
root.put("/portal/api/forwarding/rules/789", json={"target": "https://brazil.example.net/scan", "mode": "fixed",
                                                   "enabled": True}, headers=ORIGIN)
check("the longest prefix first for the resolver", seen().rule_for("01", "07891234500014")[0] == "789"
      and seen().rule_for("00", "395212345678901236")[0] == "9521234")
r = root.delete("/portal/api/forwarding/rules/789", headers=ORIGIN)
check("removed", r.status_code == 200 and code(r) == "forwarding.ruleRemoved"
      and [x["prefix"] for x in r.get_json()["rules"]] == ["9521234"] and seen().rule_for("01", "07891234500014") is None)
check("…in the journal", (store.col("journal").find_one({"action": "forwarding-rule-remove"}) or {}).get("detail")
      == "prefix=789 target=https://brazil.example.net/scan")
check("…removing it again: 404", root.delete("/portal/api/forwarding/rules/789", headers=ORIGIN).status_code == 404)

# ================================================================== the request log and the overview
EVENTS = requestlog.events()
NOW = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(minutes=2)


def event(value, o, n=1, to=None):
    for _ in range(n):
        EVENTS.insert_one({"t": NOW, "m": "GET", "f": "uri", "s": 307, "o": o, "dev": "mobile", "ms": 10, "ai": "01",
                           "v": value, "cp": value[1:], "to": to, "expiresAt": NOW + timedelta(days=90)})


event("09521234000019", "forwarded", 3, "other.example.net")
event("09521234000026", "forwarded", 1, "other.example.net")
event("07891234500014", "fallback", 2, "id.gs1.org")
event("09506000134376", "fallback", 5, "id.gs1.org")
event("07891234500014", "redirect", 4, "acme.example.com")
root.put("/portal/api/forwarding/fallback", json={"enabled": True, "url": ""}, headers=ORIGIN)
d = root.get("/portal/api/overview?days=7").get_json()
q = d["quality"]
check("the overview's quality: requests forwarded by rules and sent to the last resort, apart", q.get("forwarded") == 4
      and q.get("fallback") == 7 and q["previous"]["forwarded"] == 0 and q["answered"] == 4 and q["unknown"] == 0, q)
fw = d.get("forwarding") or {}
check("the Forwarding group for server administrators: the last resort on, its address, its requests",
      fw.get("fallback", {}).get("enabled") is True and fw["fallback"]["url"] == "https://id.gs1.org"
      and fw["fallback"]["requests"] == 7 and fw["fallback"]["previous"] == 0, fw)
check("…the codes most sent there", [(c["anchor"], c["requests"]) for c in fw["fallback"]["topCodes"]]
      == [("/01/09506000134376", 5), ("/01/07891234500014", 2)], fw["fallback"]["topCodes"])
check("…the rules on and their requests, by rule", fw["rules"]["active"] == 1 and fw["rules"]["total"] == 1
      and fw["rules"]["requests"] == 4 and fw["rules"]["byRule"] == [{"prefix": "9521234",
                                                                       "target": "https://other.example.net",
                                                                       "enabled": True, "requests": 4}], fw["rules"])
d = ana.get("/portal/api/overview?days=7").get_json()
check("an organisation's administrator: its codes sent to the last resort counted, no Forwarding group",
      d.get("forwarding") is None and d["quality"]["fallback"] == 2 and d["quality"]["forwarded"] == 0, d["quality"])
d = root.get(f"/portal/api/overview?days=7&tenant={acme['id']}").get_json()
check("…a server administrator viewing one organisation: no Forwarding group either", d.get("forwarding") is None)
labels = {"overview.group.forwarding": "Encaminhamento", "overview.f.fallbackRequests": "Enviadas ao último recurso",
          "overview.sheet.byRule": "Regras", "overview.on": "ligado"}
r = root.post("/portal/api/overview/export", json={"format": "csv", "days": 7, "labels": labels}, headers=ORIGIN)
rows = list(csv.reader(io.StringIO(r.data.decode("utf-8-sig")), delimiter=";"))
check("the export: the group's figures in the summary", ["Encaminhamento", "Enviadas ao último recurso", "7", "0"] in rows
      and ["Encaminhamento", "overview.f.fallback", "ligado", ""] in rows
      and ["Encaminhamento", "overview.f.rulesRequests", "4", "0"] in rows, [x for x in rows if x and x[0] == "Encaminhamento"])
r = root.post("/portal/api/overview/export", json={"format": "xlsx", "days": 7, "labels": labels}, headers=ORIGIN)
from openpyxl import load_workbook  # noqa: E402
book = load_workbook(io.BytesIO(r.data))
check("…and the sheets of the codes sent to the last resort and of the rules", "Regras" in book.sheetnames
      and [c.value for c in book["Regras"][2]] == ["9521234", "https://other.example.net", "ligado", 4], book.sheetnames)
r = root.get("/portal/api/requests?outcome=fallback")
check("the Requests list filters by the new outcomes", r.status_code == 200
      and len(r.get_json()["events"]) == 7 and all(e["outcome"] == "fallback" for e in r.get_json()["events"]),
      (r.status_code, r.get_json()))

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
