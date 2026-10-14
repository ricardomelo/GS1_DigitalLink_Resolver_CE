"""
The overview (Documentation/adr/0007-overview-indicators.md), through the portal's Flask test client: the period and
the previous one, usage and resolution quality from the daily aggregates and today's raw events (the same figures
whether a day is counted yet or not), the scope of each user (an organisation, a user's prefixes, the operator, the
whole installation, support access), record quality from the resolver's records, operation, governance for
administrators only, the exports (XLSX and CSV, cells protected), the tenant API's /stats and its scope, failures of
the resolver and of the log, and the texts in both languages.

Runs against mongomock (the portal's user may only read resolver_log on a real database, and this program writes
events into it, as test_requestlog.py does).

  pip install -r portal/requirements.txt mongomock
  python dev-tests/portal/test_overview.py
"""
import csv
import io
import os
import re
import sys
import tempfile
import threading
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook
from werkzeug.serving import make_server

HERE = os.path.dirname(os.path.abspath(__file__))
PORTAL = os.environ.get("PORTAL_DIR", os.path.join(HERE, "..", "..", "portal"))
CONFIG = tempfile.mkdtemp()
DATA_ENTRY_PORT = 3398
sys.path[:0] = [HERE, PORTAL]

os.environ.update(DATA_ENTRY_URL=f"http://127.0.0.1:{DATA_ENTRY_PORT}/api", SESSION_TOKEN="tok",
                  PORTAL_CONFIG_DIR=CONFIG, PORTAL_ORIGIN="https://id.example.org", PORTAL_COOKIE_SECURE="false",
                  RESOLVER_PUBLIC_URL="https://id.example.org", PORTAL_TIME_ZONE="America/Sao_Paulo")
os.environ["PORTAL_DB_URI"] = "mongomock://"
os.environ["PORTAL_DB_NAME"] = f"devtest_overview_{os.getpid()}"
os.environ["PORTAL_WORKER"] = "embedded"
for name in ("RESOLVER_LOG_DAYS", "PORTAL_STATS_DAYS", "PORTAL_MAIL_HOST"):
    os.environ.pop(name, None)

import mock_data_entry  # noqa: E402
import store  # noqa: E402
import users  # noqa: E402
users.create("root", "root-password-long", "admin", must_change=False)
import apikeys  # noqa: E402
import app as portal  # noqa: E402
import overview  # noqa: E402
import requestlog  # noqa: E402
import tenants  # noqa: E402

threading.Thread(target=make_server("127.0.0.1", DATA_ENTRY_PORT, mock_data_entry.app, threaded=True).serve_forever,
                 daemon=True).start()
failures = []
ORIGIN = {"Origin": "https://id.example.org"}


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + ("" if condition else f"  → {detail}"))
    if not condition:
        failures.append(name)


def gtin(body13: str) -> str:
    total = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(body13)))
    return body13 + str((10 - total % 10) % 10)


def client(username, password):
    c = portal.app.test_client()
    c.post("/portal/api/login", json={"username": username, "password": password}, headers=ORIGIN)
    return c


def code(r):
    return (r.get_json(silent=True) or {}).get("code")


def record(value, qualifiers=None, description="Item"):
    body = {"key": "01", "value": value, "description": description,
            "links": [{"linkType": "gs1:pip", "url": "https://example.org/x", "hreflang": ["en"], "title": "t"}]}
    if qualifiers:
        body["qualifiers"] = qualifiers
    return body


EVENTS = requestlog.events()
UTC = timezone.utc
ZONE = requestlog.zone()
TODAY = requestlog.today()


def at(days_ago: int, hour: int = 12) -> datetime:
    """Noon (or another hour) of a day of the installation's time zone, as a stored date."""
    local = datetime.combine(TODAY - timedelta(days=days_ago), datetime.min.time()).replace(hour=hour, tzinfo=ZONE)
    return min(local, datetime.now(ZONE) - timedelta(minutes=1)).astimezone(UTC).replace(tzinfo=None)


def event(value, days_ago=0, o="redirect", dev="mobile", n=1, **extra):
    """n events of a GTIN (None: an invalid request without a key) on a day."""
    for _ in range(n):
        doc = {"t": at(days_ago), "m": "GET", "f": "uri", "s": 307 if o == "redirect" else 404, "o": o, "dev": dev,
               "ms": 20}
        if value:
            doc.update(ai="01", v=value, cp=value[1:])
        doc["expiresAt"] = doc["t"] + timedelta(days=90)
        doc.update(extra)
        EVENTS.insert_one(doc)


# ================================================================== the installation
acme = tenants.create("Acme Foods", ["7891234"], "root")
users.create("ana", "ana-password-long", "admin", tenant=acme["id"], must_change=False)
users.create("leo", "leo-password-long", "reader", tenant=acme["id"], prefixes=["78912345"], must_change=False)
users.create("oto", "oto-password-long", "editor", must_change=False)
users.create("pia", "pia-password-long", "admin", tenant=acme["id"], prefixes=["78912345"], must_change=False)
root = client("root", "root-password-long")
ana, leo, oto = client("ana", "ana-password-long"), client("leo", "leo-password-long"), client("oto", "oto-password-long")

ACME_A = gtin("0789123450001")      # inside leo's 78912345: a key record, scanned
ACME_B = gtin("0789123410001")      # Acme, outside 78912345: a key record nobody scans
ACME_C = gtin("0789123450002")      # inside 78912345: only a batch, no key record
ACME_X = gtin("0789123419999")      # Acme's prefix (outside 78912345), never registered
OPER = gtin("0950600013435")        # the operator's
for c, value, quals in ((ana, ACME_A, None), (ana, ACME_B, None), (ana, ACME_C, "/10/L1"), (oto, OPER, None)):
    r = c.post("/portal/api/record", json=record(value, quals), headers=ORIGIN)
    check(f"record {value}{quals or ''} saved", r.status_code in (200, 201), (r.status_code, r.get_json()))

# Requests: two days ago (counted by the worker), today (raw), and 40 days ago (the previous 30-day period)
event(ACME_A, 2, n=5)
event(ACME_A, 1, o="noLink", n=2, lt="gs1:certificationInfo")
event(ACME_A, 0, n=3, lang="pt-br")
event(ACME_A, 0, dev="automated")
event(ACME_X, 2, o="unknown", n=4)
event(ACME_C, 0, n=1, q={"10": "L1"}, f="compressed")
event(OPER, 0, n=2, dev="desktop")
event(OPER, 0, o="compressed", dev="desktop", s=200)    # a ?compress request: answered
event(None, 0, o="invalid", dev="automated", p="/=1+1")
event(None, 0, o="invalid", s=400, p="/probe")         # a person's mistyped address: no key
event(ACME_A, 40, n=4)
event(ACME_A, 0, n=1, lang="=1+1")                 # a typed language, to prove cells are protected

# ================================================================== period and scope
r = ana.get("/portal/api/overview")
data = r.get_json()
check("GET overview answers for any user (30 days by default)", r.status_code == 200 and data["period"]["days"] == 30,
      r.status_code)
check("…the period ends today and starts 29 days before, in the installation's time zone",
      data["period"]["to"] == TODAY.isoformat() and data["period"]["from"] == (TODAY - timedelta(days=29)).isoformat()
      and data["period"]["timeZone"] == "America/Sao_Paulo", data["period"])
check("…the previous period is the 30 days before it",
      data["period"]["previousTo"] == (TODAY - timedelta(days=30)).isoformat()
      and data["period"]["previousFrom"] == (TODAY - timedelta(days=59)).isoformat(), data["period"])
check("another period: 422 overview.badPeriod", ana.get("/portal/api/overview?days=15").status_code == 422
      and code(ana.get("/portal/api/overview?days=15")) == "overview.badPeriod"
      and ana.get("/portal/api/overview?days=x").status_code == 422)

u, q = data["usage"], data["quality"]
check("usage before the worker counted: today's raw events only (Acme: 3 + 1 + 1 scans by people)",
      u["scans"] == 5 and u["automated"] == 1, (u["scans"], u["automated"]))
# the embedded worker's first housekeeping ran before these events existed: count every day again
store.col("settings").delete_many({"_id": requestlog.SETTING_FINAL})
requestlog.aggregate()
data = ana.get("/portal/api/overview").get_json()
u, q, rec = data["usage"], data["quality"], data["records"]
check("usage after the worker counted: past days from the aggregates (5 + 2 + 3 + 4 + 1 + 1 scans by people)",
      u["scans"] == 16 and u["automated"] == 1, (u["scans"], u["automated"]))
check("…the previous period from the aggregates too (the 4 scans of 40 days ago)",
      u["previous"] == {"scans": 4, "automated": 0, "pageViews": 0, "clicks": 0}, u["previous"])
check("…one entry per day, oldest first, adding up to the period's scans",
      len(u["perDay"]) == 30 and u["perDay"][-1]["day"] == TODAY.isoformat()
      and sum(d["scans"] for d in u["perDay"]) == u["scans"], [d for d in u["perDay"] if d["scans"]])
two_days = next(d for d in u["perDay"] if d["day"] == (TODAY - timedelta(days=2)).isoformat())
check("…a day counts its not-found requests (the unknown code's 4 of 9 scans)",
      two_days == {"day": two_days["day"], "scans": 9, "automated": 0, "notFound": 4}, two_days)
check("…the most scanned codes, the most first", [c["anchor"] for c in u["topCodes"]][:3]
      == [f"/01/{ACME_A}", f"/01/{ACME_X}", f"/01/{ACME_C}"] and u["topCodes"][0]["scans"] == 11, u["topCodes"])
check("…shares by device, form, language and link type",
      u["device"] == {"mobile": 15, "automated": 1, "desktop": 0} or (u["device"].get("mobile") == 16
                                                                     and u["device"].get("automated") == 1),
      u["device"])
check("…form and language counted", u["form"].get("compressed") == 1 and ["pt-br", 3] in u["language"]
      and ["gs1:certificationInfo", 2] in u["linkType"], (u["form"], u["language"], u["linkType"]))
check("quality: requests, answered, unknown and noLink (automated requests included)",
      (q["requests"], q["answered"], q["unknown"], q["noLink"], q["invalid"]) == (17, 11, 4, 2, 0),
      (q["requests"], q["answered"], q["unknown"], q["noLink"], q["invalid"]))
check("…the previous period's quality", q["previous"]["requests"] == 4 and q["previous"]["answered"] == 4, q["previous"])
check("…the average answer time", q["avgMs"] == 20 and q["p95Max"] is None, (q["avgMs"], q["p95Max"]))
check("…scanned codes nobody registered, first", q["unknownCodes"] == [
    {"key": "01", "value": ACME_X, "anchor": f"/01/{ACME_X}", "requests": 4, "registered": False}], q["unknownCodes"])
check("…codes without the link asked", [(c["anchor"], c["requests"], c["registered"]) for c in q["noLinkCodes"]]
      == [(f"/01/{ACME_A}", 2, True)], q["noLinkCodes"])

check("records: Acme's three records in three keys", rec["available"] and rec["records"] == 3 and rec["keys"] == 3
      and rec["byKey"] == {"01": 3}, rec)
check("…a batch without a key record", rec["noKeyRecord"] == 1 and rec["breaksRules"] == 0 and rec["noDescription"] == 0,
      rec)
check("…keys nobody scanned in the period", rec["notScanned"] == 1 and rec["notScannedSample"] == [f"/01/{ACME_B}"], rec)
check("…no link check yet", rec["lastCheck"] is None and rec["linkProblems"] == 0, rec)
check("operation for an organisation: the last link check, the log's retention and its alerts, no server details",
      data["operation"] == {"lastCheck": None, "logDays": 90, "alerts": []}, data["operation"])

leo_data = leo.get("/portal/api/overview").get_json()
check("a user limited to a prefix: only that prefix's requests (A and C, not B or the unknown code)",
      leo_data["usage"]["scans"] == 12 and leo_data["quality"]["unknown"] == 0, leo_data["usage"]["scans"])
check("…and only that prefix's records", leo_data["records"]["records"] == 2, leo_data["records"])
check("…no governance for a reader", leo_data["governance"] is None)

oto_data = oto.get("/portal/api/overview").get_json()
check("an operator user: the operator's requests and records only",
      oto_data["usage"]["scans"] == 3 and oto_data["usage"]["device"] == {"desktop": 3}
      and oto_data["records"]["records"] == 1 and not oto_data["whole"], oto_data["usage"])
check("…a compressed link's answer counts as answered", (oto_data["quality"]["requests"], oto_data["quality"]["answered"])
      == (3, 3), oto_data["quality"])

root_data = root.get("/portal/api/overview").get_json()
check("a server administrator: the whole installation, requests without a key included",
      root_data["whole"] and root_data["usage"]["scans"] == 20 and root_data["usage"]["automated"] == 2
      and root_data["quality"]["invalid"] == 2, (root_data["usage"]["scans"], root_data["quality"]["invalid"]))
check("…but only codes in the lists of codes", all(c["key"] == "01" and c["value"] for c in root_data["usage"]["topCodes"])
      and len(root_data["usage"]["topCodes"]) == 4, root_data["usage"]["topCodes"])
check("…the slowest day's 95th percentile (the whole installation's daily counts)",
      root_data["quality"]["p95Max"] == 20, root_data["quality"]["p95Max"])
check("…the server's operation: worker, e-mail, the last day counted",
      set(root_data["operation"]) == {"lastCheck", "logDays", "worker", "mail", "countedUntil", "alerts", "checks"}
      and root_data["operation"]["mail"] == "off" and root_data["operation"]["countedUntil"], root_data["operation"])
root.post("/portal/api/support", json={"tenant": acme["id"], "reason": "Ticket 9"}, headers=ORIGIN)
support = root.get("/portal/api/overview").get_json()
check("support access: the organisation's figures", not support["whole"] and support["usage"]["scans"] == 16
      and support["records"]["records"] == 3, support["usage"]["scans"])
root.delete("/portal/api/support", headers=ORIGIN)

week = ana.get("/portal/api/overview?days=7").get_json()
check("7 days: the same recent scans, nothing in the previous week",
      week["usage"]["scans"] == 16 and week["usage"]["previous"]["scans"] == 0 and len(week["usage"]["perDay"]) == 7,
      week["usage"])
quarter = ana.get("/portal/api/overview?days=90").get_json()
check("90 days: the scans of 40 days ago are in the period", quarter["usage"]["scans"] == 20, quarter["usage"]["scans"])

# Today's figures are the same whether the worker counted today or not
requestlog.count_day(TODAY)
again = ana.get("/portal/api/overview").get_json()
check("the same figures after the worker counted today too (today always from the raw events)",
      again["usage"]["scans"] == 16 and again["quality"]["requests"] == 17, again["usage"]["scans"])

# ================================================================== governance (administrators)
gov = data["governance"]
check("governance for an organisation's administrator: record changes by kind and by user",
      gov["changes"].get("create") == 3 and gov["byUser"] == [["ana", 3]], gov)
check("…users by role (the organisation's)", gov["users"] == {"admin": 2, "reader": 1}, gov["users"])
pia_gov = client("pia", "pia-password-long").get("/portal/api/overview").get_json()["governance"]
check("…an administrator limited to a prefix: the changes of that prefix's records only",
      pia_gov["byUser"] == [["ana", 2]] and pia_gov["changes"] == {"create": 2}, pia_gov)
check("…sign-ins counted", gov["signIns"] >= 2 and gov["signInsRefused"] == 0, gov)
key_doc = apikeys.create("Stats", ["stats:read"], 10, [], acme["id"], "ana")
apikeys.create("Records", ["records:read"], 90, [], acme["id"], "ana")
gov = ana.get("/portal/api/overview").get_json()["governance"]
check("…API keys active and expiring within 14 days", gov["apiKeys"] == {"active": 2, "expiring": 1}, gov["apiKeys"])
root_gov = root.get("/portal/api/overview").get_json()["governance"]
check("…a server administrator: every user of the installation", sum(root_gov["users"].values()) == 5, root_gov["users"])

# ================================================================== the tenant API
stats_key, records_key = key_doc[0], None
api_client = portal.app.test_client()
r = api_client.get("/portal/api/v1/stats", headers={"Authorization": f"Bearer {stats_key}"})
body = r.get_json()
check("tenant API: /stats with the scope stats:read", r.status_code == 200
      and set(body) == {"period", "usage", "quality", "records"} and body["usage"]["scans"] == 16
      and body["records"]["records"] == 3, (r.status_code, body and list(body)))
check("…no governance, no operation", "governance" not in body and "operation" not in body)
check("…a period: 7 days; another: 422", api_client.get("/portal/api/v1/stats?days=7", headers={
    "Authorization": f"Bearer {stats_key}"}).get_json()["period"]["days"] == 7
      and api_client.get("/portal/api/v1/stats?days=8", headers={"Authorization": f"Bearer {stats_key}"}).status_code == 422)
other = apikeys.create("No stats", ["records:read"], 30, [], acme["id"], "ana")[0]
r = api_client.get("/portal/api/v1/stats", headers={"Authorization": f"Bearer {other}"})
check("…a key without stats:read: 403 api.scope", r.status_code == 403 and code(r) == "api.scope", r.status_code)
check("…the scope is offered in the portal", "stats:read" in apikeys.SCOPES)
narrow = apikeys.create("Narrow stats", ["stats:read"], 30, ["78912345"], acme["id"], "ana")[0]
r = api_client.get("/portal/api/v1/stats", headers={"Authorization": f"Bearer {narrow}"})
check("…a key limited to a prefix sees that prefix only", r.status_code == 200 and r.get_json()["usage"]["scans"] == 12,
      (r.status_code, r.get_data(as_text=True)[:300]))

# ================================================================== exports
labels = {"overview.sheet.summary": "Indicadores", "overview.scans": "Leituras", "overview.group.usage": "Uso",
          "overview.sheet.topCodes": "Mais lidos: [top]"}
r = ana.post("/portal/api/overview/export", json={"format": "xlsx", "days": 30, "labels": labels}, headers=ORIGIN)
check("export: XLSX", r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
      and "portal-overview-" in r.headers.get("Content-Disposition", ""), r.status_code)
book = load_workbook(io.BytesIO(r.data))
check("…one sheet per list, named in the user's language (characters Excel refuses removed)",
      book.sheetnames[0] == "Indicadores" and "Mais lidos   top" in book.sheetnames and len(book.sheetnames) >= 6,
      book.sheetnames)
summary = [[c.value for c in row] for row in book["Indicadores"].iter_rows()]
check("…the scans with the previous period", ["Uso", "Leituras", 16, 4] in summary, summary[:3])
check("…a typed value protected against formulas",
      not any(str(v).startswith("=") for row in summary for v in row if v is not None), [r for r in summary if "1+1" in str(r)])
top = [[c.value for c in row] for row in book["Mais lidos   top"].iter_rows()]
check("…the codes scanned (up to 100)", top[1][0] == f"/01/{ACME_A}" and top[1][1] == 11, top[:3])
r = ana.post("/portal/api/overview/export", json={"format": "csv", "days": 7, "labels": {}}, headers=ORIGIN)
text = r.data.decode("utf-8")
rows = list(csv.reader(io.StringIO(text.lstrip("﻿")), delimiter=";"))
check("export: CSV with a BOM, ';' and the message codes as labels when none are sent",
      text.startswith("﻿") and rows[0] == ["overview.col.group", "overview.col.indicator", "overview.col.value",
                                                "overview.col.previous"]
      and ["overview.group.usage", "overview.scans", "16", "0"] in rows, rows[:3])
check("…cells protected", any(row[2].startswith("'=") for row in rows) and not any(row[2].startswith("=") for row in rows),
      [row for row in rows if "1+1" in row[2]])
events = store.col("journal").count_documents({"action": "overview-export"})
check("…each export is a journal event", events == 2, events)
check("…a writing call needs the portal's origin", ana.post("/portal/api/overview/export", json={"format": "csv"},
                                                             headers={"Origin": "https://evil.example"}).status_code == 403)

# ================================================================== failures
saved_resolver = portal.resolver


def down(*_a, **_k):
    raise portal.UpstreamError("upstream.unreachable", 502)


portal.resolver = down
r = ana.get("/portal/api/overview")
check("the resolver cannot be read: the overview still answers, without records",
      r.status_code == 200 and r.get_json()["records"]["available"] is False and r.get_json()["usage"]["scans"] == 16,
      r.status_code)
portal.resolver = saved_resolver
saved_events = requestlog.events


class Down:
    def find(self, *a, **k):
        from pymongo.errors import OperationFailure
        raise OperationFailure("not authorized on resolver_log", 13)
    aggregate = find


requestlog.events = lambda: Down()
r = ana.get("/portal/api/overview")
check("the log cannot be read: 503 requests.unavailable", r.status_code == 503 and code(r) == "requests.unavailable",
      r.status_code)
requestlog.events = saved_events
check("signed out: 401", portal.app.test_client().get("/portal/api/overview").status_code == 401)

# ================================================================== texts and description
i18n = open(os.path.join(PORTAL, "static", "i18n.js"), encoding="utf-8").read()
app_js = open(os.path.join(PORTAL, "static", "app.js"), encoding="utf-8").read()
html = open(os.path.join(PORTAL, "static", "index.html"), encoding="utf-8").read()
used = set(re.findall(r'"(overview\.[A-Za-z0-9.:-]+[A-Za-z0-9-])"', app_js + html))
used |= {f"overview.g.change.{a}" for a in ("create", "update", "delete", "import")}
used |= {f"overview.g.users.{r}" for r in ("admin", "editor", "reader", "disabled")}
used |= {f"overview.form.{f}" for f in ("uri", "compressed", "epc", "-")}
used |= {f"overview.q.{n}" for n in ("requests", "answered", "unknown", "noLink", "invalid", "error")}
used |= {f"overview.r.{n}" for n in ("records", "keys", "noKeyRecord", "breaksRules", "noDescription", "notScanned",
                                     "linkProblems", "byKey", "unavailable")}
used |= {"overview.badPeriod", "menu.overview", "menu.editor", "apikeys.scope.stats:read", "apikeys.scopeShort.stats:read"}
used |= set(re.findall(r'"(overview\.[A-Za-z0-9.:-]+[A-Za-z0-9-])"', open(os.path.join(PORTAL, "overview.py")).read()))
missing = sorted(k for k in used if i18n.count(f'"{k}"') != 2)
check(f"every overview text in both languages ({len(used)} keys)", not missing, missing)
spec = open(os.path.join(PORTAL, "openapi.json"), encoding="utf-8").read()
check("the API's description has /stats and the scope stats:read", '"/stats"' in spec and '"stats:read"' in spec)
login_js = open(os.path.join(PORTAL, "static", "login.js"), encoding="utf-8").read()
check("sign-in opens the overview", 'location.replace("./#overview")' in login_js)

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
