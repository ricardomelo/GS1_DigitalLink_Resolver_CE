"""
The overview in a real browser (Playwright/Chromium), the page opened after sign-in
(Documentation/adr/0007-overview-indicators.md): sign-in lands on it, its five groups of indicators for an
organisation's administrator, the period selector, a scanned code nobody registered opened in the editor, the exports,
a reader without governance, a server administrator's whole installation, both languages and a phone's width. Events
are written straight into resolver_log.requests, as the resolver would.

  pip install -r portal/requirements.txt playwright mongomock openpyxl && playwright install chromium
  python dev-tests/portal/test_overview_e2e.py               # SCREENSHOTS=<folder> keeps a picture of each step
"""
import io
import os
import sys
import tempfile
import threading
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from werkzeug.serving import make_server

HERE = os.path.dirname(os.path.abspath(__file__))
PORTAL = os.environ.get("PORTAL_DIR", os.path.join(HERE, "..", "..", "portal"))
CONFIG = tempfile.mkdtemp()
PORT, DATA_ENTRY_PORT = 8196, 3196
SHOTS = os.environ.get("SCREENSHOTS")
os.environ.update(DATA_ENTRY_URL=f"http://127.0.0.1:{DATA_ENTRY_PORT}/api", SESSION_TOKEN="tok",
                  PORTAL_CONFIG_DIR=CONFIG, PORTAL_DB_NAME=f"devtest_overview_e2e_{os.getpid()}",
                  PORTAL_ORIGIN=f"http://127.0.0.1:{PORT}", PORTAL_COOKIE_SECURE="false",
                  RESOLVER_PUBLIC_URL=f"http://127.0.0.1:{PORT}", PORTAL_TIME_ZONE="America/Sao_Paulo")
os.environ["PORTAL_DB_URI"] = "mongomock://"          # events are written into resolver_log, which the portal only reads
os.environ["PORTAL_WORKER"] = "embedded"
os.environ.pop("RESOLVER_LOG_DAYS", None)
sys.path[:0] = [HERE, PORTAL]

import mock_data_entry  # noqa: E402
import users  # noqa: E402
import tenants  # noqa: E402
import app as portal  # noqa: E402
import requestlog  # noqa: E402

acme = tenants.create("Acme Foods", ["7891234"], "root")
users.create("root", "root-password-long", "admin", must_change=False)
users.create("ana", "ana-password-long", "admin", tenant=acme["id"], must_change=False)
users.create("edu", "edu-password-long", "reader", tenant=acme["id"], must_change=False)
for application, port in [(mock_data_entry.app, DATA_ENTRY_PORT), (portal.app, PORT)]:
    threading.Thread(target=make_server("127.0.0.1", port, application, threaded=True).serve_forever, daemon=True).start()

BASE = f"http://127.0.0.1:{PORT}/portal/"
failures = []
EVENTS = requestlog.events()
NOW = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
ACME = "07891234500014"
UNKNOWN = "07891234999996"
OPER = "09506000134352"


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + ("" if condition else f"  → {detail}"))
    if not condition:
        failures.append(name)


def event(value=None, minutes=5, o="redirect", s=307, dev="mobile", **extra):
    t = NOW - timedelta(minutes=minutes)
    doc = {"t": t, "expiresAt": t + timedelta(days=90), "m": "GET", "f": "uri", "s": s, "o": o, "dev": dev, "ms": 12}
    if value:
        doc.update(ai="01", v=value, cp=value[1:])
    doc.update(extra)
    EVENTS.insert_one(doc)


def sign_in(page, user, password):
    page.goto(BASE + "login")
    page.fill("#username", user); page.fill("#password", password); page.click("#login-submit")
    page.wait_for_url(BASE + "#overview", timeout=5000)
    page.wait_for_timeout(1200)


def shot(page, name):
    if SHOTS:
        os.makedirs(SHOTS, exist_ok=True)
        page.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=True)


def texts(page, selector):
    return page.eval_on_selector_all(selector, "ns => ns.map(n => n.innerText)")


# Acme's record, scanned on several days; a code nobody registered; the operator's scans
client = portal.app.test_client()
client.post("/portal/api/login", json={"username": "ana", "password": "ana-password-long"},
            headers={"Origin": f"http://127.0.0.1:{PORT}"})
client.post("/portal/api/record", json={"key": "01", "value": ACME, "description": "Acme biscuits", "links": [
    {"linkType": "gs1:pip", "url": "https://acme.example/b", "hreflang": ["en"], "title": "Biscuits"}]},
            headers={"Origin": f"http://127.0.0.1:{PORT}"})
for minutes in (2, 3, 4, 60 * 24 * 3, 60 * 24 * 3 + 1, 60 * 24 * 10):
    event(ACME, minutes, lang="pt-br")
for minutes in (5, 6, 60 * 24 * 2):
    event(UNKNOWN, minutes, o="unknown", s=404)
event(ACME, 7, dev="automated")
event(OPER, 8, dev="desktop")
event(None, 9, o="invalid", s=400, p="/wp-login.php", dev="automated")
for day in range(1, 12):
    requestlog.count_day(requestlog.today() - timedelta(days=day))


with sync_playwright() as p:
    browser = p.chromium.launch()
    errors = []
    ctx = browser.new_context(locale="en-GB", viewport={"width": 1280, "height": 1000}, accept_downloads=True)
    page = ctx.new_page()
    page.on("pageerror", lambda e: errors.append(str(e)))

    # ------------------------------------------------------------ an organisation's administrator
    sign_in(page, "ana", "ana-password-long")
    check("sign-in opens the overview", page.is_visible("#overview-view") and not page.is_visible("#editor-view")
          and page.inner_text("#hero-title").strip() == "Overview", page.inner_text("#hero-title"))
    check("…the scope and the period are written", "Acme Foods" in page.inner_text("#overview-scope")
          and "America/Sao_Paulo" in page.inner_text("#overview-scope"), page.inner_text("#overview-scope"))
    scans = page.inner_text("#overview-tile-scans")
    check("usage: the organisation's scans by people (6 of the record, 3 of the unknown code)",
          "9" in scans.split("\n")[1], scans)
    bars = page.eval_on_selector_all("#overview-chart .requests-bar", "bs => bs.length")
    check("…one bar per day of the period", bars == 30, bars)
    check("…the bars have a text alternative", "9 scans" in (page.get_attribute("#overview-chart", "aria-label") or ""),
          page.get_attribute("#overview-chart", "aria-label"))
    top = texts(page, "#overview-top-body tr")
    check("…the most scanned codes", top and top[0].startswith("/01/" + ACME), top)
    shares = page.inner_text("#overview-shares")
    check("…shares by device and language", "Phone" in shares and "pt-br" in shares, shares[:200])
    unknown = page.inner_text("#overview-tile-unknown")
    check("quality: the scanned codes nobody registered, highlighted", "3" in unknown
          and "is-warn" in page.get_attribute("#overview-tile-unknown", "class"), unknown)
    rows = texts(page, "#overview-unknown-body tr")
    check("…listed with a Register button", rows and rows[0].startswith("/01/" + UNKNOWN) and "Register" in rows[0], rows)
    records = page.inner_text("#overview-records-tiles")
    check("records: one record, scanned", "Records\n1" in records and "Keys not scanned\n0" in records, records)
    check("record quality: links never checked, no active alerts (from Operation, ADR 0013)",
          "never checked" in page.inner_text("#overview-records-facts")
          and "Active alerts\nnone" in page.inner_text("#overview-records-facts"), page.inner_text("#overview-records-facts"))
    check("governance for an administrator: record changes and users",
          page.is_visible("#overview-governance") and "Record changes\n1" in page.inner_text("#overview-governance-tiles")
          and "ana" in page.inner_text("#overview-users-body"), page.inner_text("#overview-governance-tiles"))
    shot(page, "overview-ana")

    page.select_option("#overview-days", "7"); page.wait_for_timeout(1200)
    check("7 days: the scans of ten days ago leave the period", "8" in page.inner_text("#overview-tile-scans").split("\n")[1]
          and page.eval_on_selector_all("#overview-chart .requests-bar", "bs => bs.length") == 7,
          page.inner_text("#overview-tile-scans"))

    with page.expect_download() as info:
        page.click("#overview-xlsx")
    book = load_workbook(io.BytesIO(open(info.value.path(), "rb").read()))
    check("Excel export: sheets named in the user's language", book.sheetnames[0] == "Indicators"
          and "Most scanned" in book.sheetnames, book.sheetnames)
    summary = [[c.value for c in row] for row in book["Indicators"].iter_rows()]
    check("…the scans of the period chosen", ["Usage", "Scans", 8, 1] in summary, summary[:3])
    with page.expect_download() as info:
        page.click("#overview-csv")
    csv_text = open(info.value.path(), encoding="utf-8-sig").read()
    check("CSV export", csv_text.startswith("Group;Indicator;Value;Previous period") and "Usage;Scans;8;1" in csv_text,
          csv_text[:120])

    page.click("#overview-unknown-body button"); page.wait_for_timeout(1200)
    check("Register opens the editor on the code", page.is_visible("#editor-view")
          and page.input_value("#key-value") == UNKNOWN, page.input_value("#key-value"))

    page.click("#user-button")
    check("the menu has Overview and the link editor first", texts(page, "#user-menu-panel .menu-item")[1:3]
          == ["Overview", "Link editor"], texts(page, "#user-menu-panel .menu-item")[:4])
    page.click("#menu-overview"); page.wait_for_timeout(1200)
    check("choosing Overview closes the menu and shows it", page.is_visible("#overview-view")
          and not page.is_visible("#menu-logout"))
    page.click("#user-button"); page.click("#menu-editor"); page.wait_for_timeout(500)
    check("Link editor goes back to the editor", page.is_visible("#editor-view") and not page.is_visible("#overview-view"))

    # ------------------------------------------------------------ Portuguese
    page.goto(BASE + "#overview"); page.wait_for_timeout(1200)
    page.select_option("#locale", "pt-BR"); page.wait_for_timeout(500)
    check("in Portuguese", page.inner_text("#hero-title").strip() == "Visão geral"
          and "Leituras" in page.inner_text("#overview-usage-tiles")
          and "Cadastrar" in page.inner_text("#overview-unknown-body"), page.inner_text("#overview-usage-tiles")[:100])
    page.select_option("#locale", "en-GB"); page.wait_for_timeout(300)

    # ------------------------------------------------------------ a reader: no governance
    reader_ctx = browser.new_context(locale="en-GB", viewport={"width": 1280, "height": 1000})
    reader = reader_ctx.new_page()
    reader.on("pageerror", lambda e: errors.append(str(e)))
    sign_in(reader, "edu", "edu-password-long")
    check("a reader sees usage, quality and records but no governance", reader.is_visible("#overview-usage-tiles")
          and not reader.is_visible("#overview-governance"))
    reader_ctx.close()

    # ------------------------------------------------------------ a server administrator: the whole installation
    root_ctx = browser.new_context(locale="en-GB", viewport={"width": 1280, "height": 1000})
    root = root_ctx.new_page()
    root.on("pageerror", lambda e: errors.append(str(e)))
    sign_in(root, "root", "root-password-long")
    check("a server administrator: the whole installation", "The whole installation" in root.inner_text("#overview-scope")
          and "10" in root.inner_text("#overview-tile-scans").split("\n")[1], root.inner_text("#overview-tile-scans"))
    operation = root.inner_text("#overview-operation")
    check("…with the worker, e-mail and the daily counts", "Job worker" in operation and "E-mail" in operation
          and "Daily counts complete up to" in operation, operation)
    check("…and the state of each monitoring check", "Resolver (through the proxy)" in operation
          and "Daily backup" in operation and "HTTPS certificate" in operation, operation)
    shot(root, "overview-root")
    # an old GCP length table, as the monitoring records it (ADR 0008): its row and the active alert
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    portal.store.col("alerts").replace_one({"_id": "gcptable"}, {
        "_id": "gcptable", "state": "alert", "since": now, "checkedAt": now,
        "detail": {"date": "2026-03-01T00:00:00Z", "days": 234, "max": 90}}, upsert=True)
    root.reload(); root.wait_for_timeout(1200)
    operation = root.inner_text("#overview-operation")
    check("…an old GCP length table: its check with the list's date and age, and the active alert",
          "list of 01/03/2026 (234 days old)" in operation and "GCP length table 234 days old" in operation, operation)
    portal.store.col("alerts").delete_one({"_id": "gcptable"})
    root_ctx.close()

    # ------------------------------------------------------------ a phone
    phone_ctx = browser.new_context(locale="en-GB", viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
    phone = phone_ctx.new_page()
    phone.on("pageerror", lambda e: errors.append(str(e)))
    sign_in(phone, "ana", "ana-password-long")
    width = phone.evaluate("document.documentElement.scrollWidth")
    check("a phone: no sideways scrolling", width <= 390, width)
    check("…tiles in one or two columns", phone.eval_on_selector("#overview-usage-tiles",
          "e => getComputedStyle(e).gridTemplateColumns.split(' ').length") <= 2)
    shot(phone, "overview-phone")
    phone_ctx.close()

    check("no errors in the page", not errors, errors)
    browser.close()

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
