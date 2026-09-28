"""
Unit tests for portal/sheet.py (spreadsheet import/export), without a browser or a resolver.

  pip install -r portal/requirements.txt
  python dev-tests/portal/test_sheet.py
"""
import io
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.environ.get("PORTAL_DIR", os.path.join(HERE, "..", "..", "portal")))

from openpyxl import load_workbook  # noqa: E402

import sheet  # noqa: E402
from gs1 import ValidationError  # noqa: E402

failures = []
PT = {"gtin": ["GTIN"], "lot": ["Lote", "Batch/lot"], "description": ["Descrição", "Description"],
      "linkType": ["Tipo de link", "Link type"], "url": ["URL"], "language": ["Idioma", "Language"],
      "title": ["Título", "Title"], "default": ["Principal", "Default"],
      "forward": ["Repassar parâmetros", "Forward query string"]}


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + ("" if condition else f"  → {detail}"))
    if not condition:
        failures.append(name)


def error_code(callable_):
    try:
        callable_()
    except ValidationError as exc:
        return exc.code
    return None


# Brazilian Excel: ";" separator, Windows-1252, Portuguese headers without the optional columns' order
csv_pt = ("Descrição;GTIN;Lote;Tipo de link;URL;Idioma;Principal;Repassar parâmetros\r\n"
          "Café torrado;7898357410015;;gs1:pip;www.exemplo.com.br/cafe;pt;sim;sim\r\n"
          "Café torrado;7898357410015;;instructions;https://exemplo.com.br/manual.pdf;pt, en;;não\r\n"
          ";;;;;;;\r\n"
          "Café torrado;7898357410015;L1;gs1:pip;https://exemplo.com.br/l1;pt;;\r\n").encode("cp1252")
rows = sheet.read_table("links.csv", csv_pt)
records, errors = sheet.parse_rows(rows, PT)
check("CSV cp1252 with ';' and Portuguese headers", len(records) == 2 and not errors, (records, errors))
product = records[0]
check("rows of the same GTIN and lot form one record", product["rows"] == [2, 3] and records[1]["rows"] == [5])
check("scheme added to bare addresses", product["links"][0]["url"] == "https://www.exemplo.com.br/cafe")
check("'gs1:' added to bare link types", product["links"][1]["linkType"] == "gs1:instructions")
check("several languages in one cell", product["links"][1]["hreflang"] == ["pt", "en"])
check("sim/não and empty defaults", product["links"][0]["default"] and not product["links"][1]["default"]
      and product["links"][0]["forward"] and not product["links"][1]["forward"] and records[1]["links"][0]["forward"])

check("canonical keys accepted without aliases", error_code(lambda: sheet.parse_rows(sheet.read_table("a.csv", b"gtin,description,linkType,url\n1,a,b,c\n"), {})) is None)
rows = sheet.read_table("x.csv", "GTIN;Descrição;Tipo de link;URL\r\n7,89836E+12;Café;gs1:pip;https://x.org\r\n".encode())
_, errors = sheet.parse_rows(rows, PT)
check("GTIN in scientific notation reported", errors and errors[0]["code"] == "import.gtinScientific", errors)
rows = sheet.read_table("x.csv", "GTIN;Descrição;Tipo de link;URL;Principal\r\n7898357410015;A;gs1:pip;https://x.org;talvez\r\n".encode())
_, errors = sheet.parse_rows(rows, PT)
check("unknown yes/no value reported", errors and errors[0]["code"] == "import.yesNo", errors)
rows = sheet.read_table("x.csv", "GTIN;Descrição;Tipo de link;URL\r\n7898357410015;A;gs1:pip;https://x.org\r\n0007898357410015;B;gs1:pip;https://y.org\r\n".encode())
records, errors = sheet.parse_rows(rows, PT)
check("leading zeros do not split a record; different descriptions reported",
      len(records) == 1 and errors and errors[0]["code"] == "import.descriptionConflict", (records, errors))
check("missing required columns", error_code(lambda: sheet.parse_rows([["GTIN", "URL"]], PT)) == "import.missingColumns")
check("empty file", error_code(lambda: sheet.parse_rows([[""], []], PT)) == "import.empty")
check("file too big", error_code(lambda: sheet.read_table("a.csv", b"x" * (sheet.MAX_FILE_BYTES + 1))) == "import.tooBig")
check("unknown format", error_code(lambda: sheet.read_table("a.pdf", b"%PDF")) == "import.format")
check("broken XLSX", error_code(lambda: sheet.read_table("a.xlsx", b"PK\x03\x04broken")) == "import.unreadable")

# Export → XLSX → import round trip
records = [{"gtin": "07898357410015", "lot": None, "description": "Café torrado", "defaultLinkType": "gs1:pip",
            "links": [{"linktype": "gs1:instructions", "href": "https://x.org/m.pdf", "hreflang": ["pt", "en"], "fwqs": False},
                      {"linktype": "gs1:pip", "href": "https://x.org/cafe", "hreflang": ["pt"], "title": "Café"}]}]
labels = {"headers": {"gtin": "GTIN", "lot": "Lote", "description": "Descrição", "linkType": "Tipo de link", "url": "URL",
                      "language": "Idioma", "title": "Título", "default": "Principal", "forward": "Repassar parâmetros"},
          "yes": "sim", "no": "não", "sheets": {"links": "Links", "linkTypes": "Tipos de link", "languages": "Idiomas"}}
rows = sheet.export_rows(records)
check("export: default link first and marked", rows[0][3] == "gs1:pip" and rows[0][7] == "yes" and rows[1][7] == "")
data = sheet.write_xlsx([list(r) for r in rows], labels, [("gs1:pip", "Página do produto", "…")], [("pt", "português")])
book = load_workbook(io.BytesIO(data))
links = book["Links"]
check("XLSX: localised headers and sheets", [c.value for c in links[1]][:4] == ["GTIN", "Lote", "Descrição", "Tipo de link"]
      and book.sheetnames == ["Links", "Tipos de link", "Idiomas"], book.sheetnames)
check("XLSX: GTIN stored as text with leading zero", links["A2"].value == "07898357410015" and links["A2"].number_format == "@")
back, errors = sheet.parse_rows(sheet.read_table("x.xlsx", data), PT)
check("XLSX round trip", not errors and len(back) == 1 and back[0]["links"][0]["default"]
      and back[0]["links"][1]["forward"] is False and back[0]["links"][1]["hreflang"] == ["pt", "en"], (back, errors))
csv_bytes = sheet.write_csv([list(r) for r in rows], labels)
check("CSV export: BOM, ';' and localised yes/no", csv_bytes.startswith(b"\xef\xbb\xbf") and b";sim;sim" in csv_bytes
      and "não".encode() in csv_bytes)
back, errors = sheet.parse_rows(sheet.read_table("x.csv", csv_bytes), PT, ["sim"], ["não"])
check("CSV round trip", not errors and back[0]["links"][1]["forward"] is False, (back, errors))

# Language tags and lots stored by other tools
import gs1  # noqa: E402
check("BCP 47 tags accepted in canonical case",
      [gs1.normalise_language(t) for t in ["pt", "pt-br", "en_US", "vi", "und", "zh-hant-tw", "es-419"]]
      == ["pt", "pt-BR", "en-US", "vi", "und", "zh-Hant-TW", "es-419"])
check("malformed tags refused", not any(gs1.normalise_language(t) for t in ["", "portuguese", "p", "en-", "en-US-x"]))
rows = sheet.read_table("x.csv", "GTIN;Descrição;Tipo de link;URL;Idioma\r\n7898357410015;A;gs1:pip;https://x.org;en-US, vi\r\n".encode())
records, _ = sheet.parse_rows(rows, PT)
check("language case kept from the file", records[0]["links"][0]["hreflang"] == ["en-US", "vi"], records[0]["links"][0])
check("template and unusual lots are not editable",
      gs1.is_editable_lot("L2026A") and not gs1.is_editable_lot("{lotnumber}") and not gs1.is_editable_lot("A/B"))

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
