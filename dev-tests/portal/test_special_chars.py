"""
Special characters in the portal: identifiers limited to ASCII (other scripts' digits refused, never an
internal error), invisible characters removed from identifiers, free text in Unicode normal form C without
control characters, addresses without spaces, line breaks, backslashes or invisible characters, spreadsheet
exports that never run as formulas (XLSX text cells, CSV apostrophe) and read back unchanged, UTF-16 text
files, and user names typed at failed sign-ins made harmless for the log and the audit CSV.

  pip install -r portal/requirements.txt
  python dev-tests/portal/test_special_chars.py
"""
import csv
import io
import os
import sys
import tempfile
import threading
import unicodedata

from openpyxl import load_workbook
from werkzeug.serving import make_server

HERE = os.path.dirname(os.path.abspath(__file__))
PORTAL = os.environ.get("PORTAL_DIR", os.path.join(HERE, "..", "..", "portal"))
CONFIG = tempfile.mkdtemp()
DATA_ENTRY_PORT = 3398
os.environ.update(DATA_ENTRY_URL=f"http://127.0.0.1:{DATA_ENTRY_PORT}/api", SESSION_TOKEN="tok",
                  PORTAL_USERS_FILE=f"{CONFIG}/users.json", PORTAL_CONFIG_DIR=CONFIG,
                  PORTAL_ORIGIN="http://localhost", PORTAL_COOKIE_SECURE="false",
                  RESOLVER_PUBLIC_URL="https://id.example.org")
sys.path[:0] = [HERE, PORTAL]

import mock_data_entry  # noqa: E402
import app as portal  # noqa: E402
import gs1  # noqa: E402
import sheet  # noqa: E402
import users  # noqa: E402
from gs1 import ValidationError  # noqa: E402

threading.Thread(target=make_server("127.0.0.1", DATA_ENTRY_PORT, mock_data_entry.app, threaded=True).serve_forever,
                 daemon=True).start()
failures = []


def check(name, condition, detail=""):
    print(("PASS " if condition else "FAIL ") + name + ("" if condition else f"  → {detail}"))
    if not condition:
        failures.append(name)


def outcome(callable_):
    """The value, the ValidationError code, or the name of any other exception (an internal error)."""
    try:
        return callable_()
    except ValidationError as exc:
        return exc.code
    except Exception as exc:  # noqa: BLE001
        return "CRASH " + type(exc).__name__


# ------------------------------------------------------------------ identifiers: ASCII only
FULL_WIDTH = "７８９８３５７４１００１５"          # 7898357410015 in full-width digits
ARABIC = "٧٨٩٨٣٥٧٤١٠٠١٥"                     # the same in Arabic-Indic digits
for ai, value in [("01", FULL_WIDTH), ("01", ARABIC), ("01", "789835741001²"), ("414", FULL_WIDTH),
                  ("00", "0" * 16 + "²0"), ("255", "²" * 13), ("8006", "²" * 18),
                  ("253", "²" * 13), ("8003", "0" + "²" * 13), ("8017", "²" * 18)]:
    result = outcome(lambda: gs1.normalise_key(ai, value))
    check(f"({ai}) {value!r}: other scripts' digits refused, no internal error",
          result in ("gtin.digitsOnly", "key.digitsOnly", "key.baseDigits"), result)
check("GIAI with other scripts' digits refused", outcome(lambda: gs1.normalise_key("8004", "٧٨٩٨ABC")) == "key.chars")
for q, key, value in [("8019", "8018", "٣"), ("8011", "8010", "٣٣"), ("7040", "417", "٣ABC"), ("10", "01", "LÓTE"),
                      ("10", "01", "A/B"), ("21", "01", "S 1")]:
    check(f"qualifier ({q}) {value!r} refused", outcome(lambda: gs1.normalise_qualifiers(key, {q: value})) == "qualifier.invalid")
check("language tag with other scripts' digits refused", gs1.normalise_language("es-٤١٩") is None)
check("GS1 Company Prefix of a user: ASCII digits only", not users.PREFIX.match("７８９８３５７")
      and users.PREFIX.match("7898357"))

# ------------------------------------------------------------------ identifiers: invisible characters removed
check("GTIN with zero-width space, BOM and no-break space accepted",
      gs1.normalise_key("01", "\u200b7898357410015\ufeff\u00a0") == "07898357410015")
check("GLN with a word joiner and a soft hyphen accepted", gs1.normalise_key("414", "95060\u2060001\u00ad34376") == "9506000134376")
check("qualifier with invisible characters accepted without them",
      gs1.normalise_qualifiers("01", {"10": "\u200bL2026A\u200e"}) == [("10", "L2026A")])
check("language tag with a BOM accepted", gs1.normalise_language("\ufeffpt-br") == "pt-BR")

# ------------------------------------------------------------------ free text
decomposed = unicodedata.normalize("NFD", "Café Açaí")
check("free text stored in normal form C", gs1.clean_text(decomposed) == "Café Açaí"
      and unicodedata.is_normalized("NFC", gs1.clean_text(decomposed)))
check("line breaks, tabs and control characters become one space",
      gs1.clean_text(" Café\r\ntorrado\t500 g\x00\u2028moído ") == "Café torrado 500 g moído")
family = "Família \U0001F468\u200d\U0001F469\u200d\U0001F467"
check("emoji sequences and other visible characters kept", gs1.clean_text(family) == family)

# ------------------------------------------------------------------ addresses
for url in ["https://a.com/x\ny", "https://a.com/x\ty", "https://a.com/x\u00a0y", "https://a.com/x y",
            "https://a.com/\u200bx", "https://a.com\\@evil.example/", "https://a.com/x\x7f"]:
    check(f"address {url!r} refused as having unusable characters",
          outcome(lambda: gs1.normalise_url(url, 1)) == "link.urlChars")
for url in ["https://exemplo.com.br/café?x=ç", "https://açúcar.com.br/", "https://a.com/busca?q=%C3%A7&x=1#frag",
            "https://a.com/(1)/'a'/\"b\""]:
    check(f"address {url!r} accepted as typed", gs1.normalise_url(url, 1) == url)

# ------------------------------------------------------------------ spreadsheets
tricky = ["=HYPERLINK(\"https://evil.example\",\"x\")", "+55 11 3068-6229", "-10% desconto", "@SUM(1)",
          "'s-Hertogenbosch", "Café ☕ <b>&amp;</b>"]
records = [{"key": "01", "value": "07898357410015", "qualifiers": [], "description": tricky[0], "defaultLinkType": "gs1:pip",
            "links": [{"linktype": code, "href": f"https://example.org/{i}", "hreflang": ["pt"], "title": title, "fwqs": True}
                      for i, (code, title) in enumerate(zip(["gs1:pip", "gs1:support", "gs1:faqs", "gs1:review",
                                                              "gs1:promotion", "gs1:recipeInfo"], tricky))]}]
PT = {"description": ["Descrição"], "linkType": ["Tipo de link"], "title": ["Título"]}

xlsx = sheet.write_xlsx(sheet.export_rows(records), {}, [("gs1:pip", "=x", "-y")], [("pt", "+z")])
book = load_workbook(io.BytesIO(xlsx))
formulas = [c.coordinate for ws in book.worksheets for line in ws.iter_rows() for c in line if c.data_type == "f"]
check("XLSX export: no cell is a formula", not formulas, formulas)
check("XLSX export: description written exactly", book.active["D2"].value == tricky[0], book.active["D2"].value)
back, errors = sheet.parse_rows(sheet.read_table("x.xlsx", xlsx), PT)
check("XLSX round trip: description and titles unchanged",
      not errors and back[0]["description"] == tricky[0] and [l["title"] for l in back[0]["links"]] == tricky, (back, errors))

csv_bytes = sheet.write_csv(sheet.export_rows(records), {})
cells = [c for row in csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig")), delimiter=";") for c in row]
risky = [c for c in cells if c.startswith(sheet.FORMULA_START)]
check("CSV export: no cell starts with = + - @", not risky, risky)
check("CSV export: such cells get an apostrophe", "'=HYPERLINK(\"https://evil.example\",\"x\")" in cells and "'+55 11 3068-6229" in cells)
back, errors = sheet.parse_rows(sheet.read_table("x.csv", csv_bytes), PT)
check("CSV round trip: description and titles unchanged, apostrophe of 's-Hertogenbosch kept",
      not errors and back[0]["description"] == tricky[0] and [l["title"] for l in back[0]["links"]] == tricky, (back, errors))

unicode_text = "GTIN\tDescrição\tTipo de link\tURL\r\n7898357410015\tAçaí ☕\tgs1:pip\thttps://x.org\r\n"
for name, data in [("Excel 'Unicode text' (UTF-16 LE)", unicode_text.encode("utf-16")),
                   ("UTF-16 BE", b"\xfe\xff" + unicode_text.encode("utf-16-be"))]:
    back, errors = sheet.parse_rows(sheet.read_table("x.txt", data), PT)
    check(f"{name} read", not errors and back and back[0]["description"] == "Açaí ☕", (back, errors))
rows = sheet.read_table("x.csv", "GTIN;Descrição;Tipo de link;URL\r\n7898357410015;A;gs1:pip;https://x.org\r\n"
                        "\u200b7898357410015\u00a0;A;gs1:support;https://y.org\r\n".encode())
back, errors = sheet.parse_rows(rows, PT)
check("invisible characters do not split a record", len(back) == 1 and len(back[0]["links"]) == 2, back)

# ------------------------------------------------------------------ through the portal's API
users.create("root", "root-password-long", "admin", must_change=False)
admin = portal.app.test_client()
check("administrator signs in", admin.post("/portal/api/login", json={"username": "root", "password": "root-password-long"}).status_code == 200)


def save(value, description="Item", url="https://example.org/x", title="t"):
    return admin.post("/portal/api/record", json={"key": "01", "value": value, "description": description,
                                                   "links": [{"linkType": "gs1:pip", "url": url, "hreflang": ["pt"], "title": title}]})


r = save("789835741001²")
check("superscript digit in a GTIN: 422 with a message, not 500", r.status_code == 422 and r.get_json()["code"] == "gtin.digitsOnly",
      (r.status_code, r.get_json()))
r = save(FULL_WIDTH)
check("full-width GTIN refused", r.status_code == 422 and r.get_json()["code"] == "gtin.digitsOnly", r.get_json())
r = save("7898357410015", url="https://example.org/a\nb")
check("address with a line break refused with its own message", r.status_code == 422 and r.get_json()["code"] == "link.urlChars"
      and r.get_json()["params"]["url"] == "https://example.org/a\\nb", r.get_json())
r = save("\u200b7898357410015", description=unicodedata.normalize("NFD", "Café\r\ntorrado"), title="Informac\u0327o\u0303es\t")
check("record saved", r.status_code in (200, 201), r.get_json())
stored = mock_data_entry.v3("01_07898357410015")
check("stored under the clean GTIN, description and title in NFC without line breaks",
      stored and stored[0]["itemDescription"] == "Café torrado" and stored[0]["links"][0]["title"] == "Informações", stored)

anonymous = portal.app.test_client()
for typed in ["=cmd|' /C calc'!A0", "evil\nuser=root action=login", "x" * 5000]:
    anonymous.post("/portal/api/login", json={"username": typed, "password": "wrong"})
audit_csv = admin.get("/portal/api/audit.csv").get_data().decode("utf-8-sig")
names = [row[1] for row in csv.reader(io.StringIO(audit_csv), delimiter=";")][1:]
check("audit CSV: a typed name starting with = does not become a formula", "'=cmd|' /C calc'!A0" in names, names)
check("audit trail: line breaks of typed names replaced", "evil?user=root action=login" in names, names)
check("audit trail: typed names cut at 64 characters", "x" * 64 + "…" in names, [len(n) for n in names])
check("audit CSV: no cell starts with = + - @",
      not [c for row in csv.reader(io.StringIO(audit_csv), delimiter=";") for c in row if c.startswith(sheet.FORMULA_START)])
check("shown_username: empty is '-'", portal.shown_username("") == "-")

print(f"\n{len(failures)} failure(s)")
sys.exit(1 if failures else 0)
