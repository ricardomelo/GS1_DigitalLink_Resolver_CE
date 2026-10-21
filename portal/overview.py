"""
The overview: the indicators of the page opened after sign-in (Documentation/adr/0007-overview-indicators.md).

Every figure is limited to the scope of the request (the scope filter of the request log, the records the user may
see, the audit trail of the scope) and computed on request from what the portal already keeps:

  usage, quality   the daily aggregates and today's raw events of the request log (requestlog.collect());
  records          the records of the record list (app.visible_records()) and the last check of every link;
  operation        server administrators only (ADR 0013): the worker, e-mail, the monitoring checks and the log's
                   counts; the last check of every link, the log's days and the scope's alerts are with records;
  governance       the journal, the user accounts, the API keys and the tenant API's calls (administrators).

app.py gathers the scope and calls build(); export_rows() and write_xlsx() turn the answer into the files.
"""
import csv
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import apikeys
import journal
import requestlog
import sheet
import store
import users

PERIODS = (7, 30, 90)
DEFAULT_DAYS = 30
TOP_SCREEN = 10
TOP_EXPORT = 100
EXPIRING_DAYS = 14
NOT_SCANNED_SAMPLE = 20
ANSWERED = ("redirect", "choice", "linkset", "compressed")
FOLLOW_UPS = ("page", "click")          # the item's page and its buttons: after a scan, not scans (ADR 0011)
REFUSED_SIGN_INS = ("login-failed", "login-refused")


class PeriodError(ValueError):
    """The period asked is not one of PERIODS."""


# ----------------------------------------------------------------------------------------------- period
def period(days) -> dict:
    """The period of the last `days` days, today included, and the one of the same length just before it."""
    try:
        days = int(days or DEFAULT_DAYS)
    except (TypeError, ValueError):
        raise PeriodError(days)
    if days not in PERIODS:
        raise PeriodError(days)
    last = requestlog.today()
    first = last - timedelta(days=days - 1)
    return {"days": days, "from": first, "to": last, "previousFrom": first - timedelta(days=days),
            "previousTo": first - timedelta(days=1), "timeZone": str(requestlog.zone())}


def _iso_period(per: dict) -> dict:
    return {k: v.isoformat() if isinstance(v, date) else v for k, v in per.items()}


# ----------------------------------------------------------------------------------------------- usage and quality
def _follow_ups(entry: dict) -> int:
    return sum(entry["outcome"].get(o, 0) for o in FOLLOW_UPS)


def _scans(entry: dict) -> int:
    """People's requests that are not follow-ups (the resolver gives automated clients' follow-ups plain outcomes)."""
    return max(0, entry["total"] - entry["automated"] - _follow_ups(entry))


def _not_found(entry: dict) -> int:
    return entry["outcome"].get("unknown", 0) + entry["outcome"].get("noLink", 0)


def _sorted_pairs(counter: dict, top: int) -> list[list]:
    return [[name, n] for name, n in sorted(counter.items(), key=lambda kv: (-kv[1], str(kv[0])))[:top]]


def _quality_counts(entry: dict) -> dict:
    outcome = entry["outcome"]
    # forwarded and fallback: sent to another resolver, by a prefix rule or as the last resort (ADR 0018)
    return {"requests": entry["total"] - _follow_ups(entry), "answered": sum(outcome.get(o, 0) for o in ANSWERED),
            "unknown": outcome.get("unknown", 0), "noLink": outcome.get("noLink", 0),
            "invalid": outcome.get("invalid", 0), "error": outcome.get("error", 0),
            "forwarded": outcome.get("forwarded", 0), "fallback": outcome.get("fallback", 0)}


def usage_and_quality(scope: dict, whole: bool, per: dict, top: int,
                      registered: set | None) -> tuple[dict, dict, dict, tuple[dict, dict]]:
    """(usage, quality, requests per key, (the period's totals, the previous period's)) of a period. registered: the
    anchors of the records in the scope, to tell a scanned code nobody registered from one whose key has records but
    none for the qualifiers scanned (None when the records could not be read)."""
    now = requestlog.collect(scope, whole, per["from"], per["to"])
    before = requestlog.collect(scope, whole, per["previousFrom"], per["previousTo"], keys=False)
    total, keys = now["total"], now["keys"]
    per_day = []
    day = per["from"]
    while day <= per["to"]:
        entry = now["days"].get(day.isoformat()) or requestlog._empty()
        scans = _scans(entry)
        per_day.append({"day": day.isoformat(), "scans": scans, "automated": entry["automated"],
                        "notFound": min(_not_found(entry), scans)})
        day += timedelta(days=1)

    def code(ai, value, entry, **extra):
        return {"key": ai, "value": value, "anchor": f"/{ai}/{value}", **extra}

    scanned = sorted(((k, e) for k, e in keys.items() if _scans(e)), key=lambda kv: (-_scans(kv[1]), str(kv[0])))
    usage = {
        "scans": _scans(total), "automated": total["automated"],
        "pageViews": total["outcome"].get("page", 0), "clicks": total["outcome"].get("click", 0),
        "previous": {"scans": _scans(before["total"]), "automated": before["total"]["automated"],
                     "pageViews": before["total"]["outcome"].get("page", 0),
                     "clicks": before["total"]["outcome"].get("click", 0)},
        "perDay": per_day,
        "topCodes": [code(ai, v, e, scans=_scans(e), notFound=min(_not_found(e), _scans(e)))
                     for (ai, v), e in scanned[:top]],
        "device": dict(total["device"]), "form": {f: n for f, n in total["form"].items() if f != "landing"},
        "language": _sorted_pairs(total["language"], top), "linkType": _sorted_pairs(total["linkType"], top),
        # the browser family and the approximate location (ADR 0006, section 8); a place is "country|region|city"
        "browser": _sorted_pairs(total["browser"], top), "country": _sorted_pairs(total["country"], top),
        "place": _sorted_pairs(total["place"], top),
    }

    def worst(outcome):
        found = sorted(((k, e) for k, e in keys.items() if e["outcome"].get(outcome)),
                       key=lambda kv: (-kv[1]["outcome"][outcome], str(kv[0])))
        return [code(ai, v, e, requests=e["outcome"][outcome],
                     registered=None if registered is None else f"/{ai}/{v}" in registered)
                for (ai, v), e in found[:top]]

    quality = {**_quality_counts(total), "previous": _quality_counts(before["total"]),
               "avgMs": round(total["ms"] / total["total"]) if total["total"] else None,
               "p95Max": max(now["p95"]) if whole and now["p95"] else None,
               "unknownCodes": worst("unknown"), "noLinkCodes": worst("noLink")}
    return usage, quality, keys, (total, before["total"])


# ----------------------------------------------------------------------------------------------- records
def record_quality(records: list[dict] | None, requested: dict, last_check: dict, top: int) -> dict:
    """The quality of the records in the scope. records: as app.visible_records() lists them (None: the resolver
    could not be read); requested: requests per (key, value) in the period; last_check: the last check of every
    link, limited to the scope ({"checkedAt", "records": {record: problems}, "urls"})."""
    check = ({"at": last_check.get("checkedAt"), "records": len(last_check.get("records") or {}),
              "urls": last_check.get("urls")} if last_check else None)
    if records is None:
        return {"available": False, "lastCheck": check, "linkProblems": check["records"] if check else 0}
    managed = [r for r in records if r.get("kind") != "other"]
    anchors = sorted({r["anchor"] for r in managed})
    by_key: dict[str, int] = {}
    for r in managed:
        by_key[r["key"]] = by_key.get(r["key"], 0) + 1
    quiet = [a for a in anchors if not requested.get(tuple(a.strip("/").split("/", 1)))]
    return {"available": True, "records": len(managed), "keys": len(anchors), "byKey": by_key,
            "noKeyRecord": sum(1 for r in managed if r.get("noKeyRecord")),
            "breaksRules": sum(1 for r in managed if r.get("breaksRules")),
            "noDescription": sum(1 for r in managed if not (r.get("description") or "").strip()),
            "notScanned": len(quiet), "notScannedSample": quiet[:top],
            "linkProblems": check["records"] if check else 0, "lastCheck": check}


# ----------------------------------------------------------------------------------------------- governance
def governance(audit_scope: dict, allowed, per: dict, tenant: str | None, whole: bool, top: int) -> dict:
    """Changes, sign-ins, users, API keys and the tenant API's calls of the scope in the period."""
    start = requestlog.day_bounds(per["from"])[0]
    end = requestlog.day_bounds(per["to"])[1]
    actions, by_user = journal.counts(audit_scope, start, end, allowed)
    accounts = [a for a in users.load().values() if whole or a.get("tenant") == tenant]
    roles: dict[str, int] = {}
    for account in accounts:
        name = "disabled" if account.get("disabled") else account.get("role", "reader")
        roles[name] = roles.get(name, 0) + 1
    now = store.now()
    keys = list(store.col("api_keys").find({})) if whole else apikeys.load(tenant)
    active = [k for k in keys if apikeys.state(k, now) == "active"]
    calls_query = {"t": {"$gte": start, "$lt": end}}
    if not whole:
        calls_query["tenant"] = tenant if tenant else {"$exists": False}
    calls = requestlog.api_calls()
    return {"changes": {a: n for a, n in actions.items() if a in journal.CATEGORIES["record"]},
            "byUser": _sorted_pairs(by_user, top),
            "signIns": actions.get("login", 0), "signInsRefused": sum(actions.get(a, 0) for a in REFUSED_SIGN_INS),
            "exports": sum(n for a, n in actions.items() if a.endswith("export")),
            "users": roles,
            "apiKeys": {"active": len(active),
                        "expiring": sum(1 for k in active if k["expiresAt"] <= now + timedelta(days=EXPIRING_DAYS))},
            "apiCalls": {"calls": calls.count_documents(calls_query),
                         "refused": calls.count_documents({**calls_query, "s": {"$gte": 400}})}}


# ----------------------------------------------------------------------------------------------- the whole answer
def build(per: dict, scope: dict, whole: bool, records: list[dict] | None, last_check: dict,
          operation: dict, governance_of=None, top: int = TOP_SCREEN, log_days: int | None = None,
          forwarding_of=None) -> dict:
    """The overview of a period (ADR 0007, section 3.2). governance_of: a function of (period, top) giving the
    governance group, or None when the user is not an administrator. log_days: how long the scope's requests are kept
    (ADR 0006, section 8; the installation's when not given). forwarding_of: a function of (requests per key, the
    period's totals, the previous period's, top) giving the Forwarding group (server administrators; ADR 0018)."""
    registered = {r["anchor"] for r in records} if records is not None else None
    usage, quality, requested, (total, before) = usage_and_quality(scope, whole, per, top, registered)
    return {"period": _iso_period(per), "whole": whole, "logDays": requestlog.global_days() if log_days is None else log_days,
            "usage": usage, "quality": quality,
            "records": record_quality(records, requested, last_check, max(top, NOT_SCANNED_SAMPLE)),
            "operation": operation,
            "governance": governance_of(per, top) if governance_of else None,
            "forwarding": forwarding_of(requested, total, before, top) if forwarding_of else None}


# ----------------------------------------------------------------------------------------------- files
def place_text(place: str) -> str:
    """"BR|Sao Paulo|Campinas" → "Campinas, Sao Paulo, BR"."""
    country, region, city = (place.split("|") + ["", ""])[:3]
    return ", ".join(part for part in (city, region, country) if part)


def _label(labels: dict, code: str) -> str:
    return labels.get(code) or code


def summary_rows(data: dict, labels: dict) -> list[list]:
    """[group, indicator, value, previous period] for every single figure of the overview."""
    rows = []

    def add(group, code, value, previous=None):
        rows.append([_label(labels, "overview.group." + group), _label(labels, code),
                     "" if value is None else value, "" if previous is None else previous])

    u, q, r = data["usage"], data["quality"], data["records"]
    add("usage", "overview.scans", u["scans"], u["previous"]["scans"])
    add("usage", "overview.automated", u["automated"], u["previous"]["automated"])
    add("usage", "overview.pageViews", u["pageViews"], u["previous"]["pageViews"])
    add("usage", "overview.clicks", u["clicks"], u["previous"]["clicks"])
    for name, n in sorted(u["device"].items()):
        add("usage", "requests.device." + name, n)
    for name, n in sorted(u["form"].items()):
        add("usage", "overview.form." + name, n)
    for name, n in u["language"]:
        add("usage", "overview.language", f"{name}: {n}")
    for name, n in u["linkType"]:
        add("usage", "overview.linkType", f"{name}: {n}")
    for name, n in u.get("browser", []):
        add("usage", "overview.browser", f"{_label(labels, 'requests.browser.' + name)}: {n}")
    for name, n in u.get("country", []):
        add("usage", "overview.country", f"{name}: {n}")
    for name, n in u.get("place", []):
        add("usage", "overview.place", f"{place_text(name)}: {n}")
    for name in ("requests", "answered", "unknown", "noLink", "invalid", "error", "forwarded", "fallback"):
        add("quality", "overview.q." + name, q[name], q["previous"][name])
    add("quality", "overview.avgMs", q["avgMs"])
    if q.get("p95Max") is not None:
        add("quality", "overview.p95Max", q["p95Max"])
    if r.get("available"):
        for name in ("records", "keys", "noKeyRecord", "breaksRules", "noDescription", "notScanned", "linkProblems"):
            add("records", "overview.r." + name, r[name])
        for ai, n in sorted(r["byKey"].items()):
            add("records", "overview.r.byKey", f"({ai}): {n}")
    else:
        add("records", "overview.r.unavailable", "")
    if r.get("lastCheck"):
        add("records", "overview.lastCheck", r["lastCheck"]["at"])
    add("records", "overview.op.logDays", data.get("logDays"))
    if r.get("alerts") is not None:
        add("records", "overview.op.alerts", len(r["alerts"]))
    op = data.get("operation") or {}       # server administrators only (ADR 0013)
    for name in ("worker", "mail", "countedUntil"):
        if name in op:
            add("operation", "overview.op." + name, op[name])
    if "alerts" in op:
        add("operation", "overview.op.alerts", len(op["alerts"]))
    fw = data.get("forwarding")            # server administrators, the whole installation (ADR 0018)
    if fw:
        add("forwarding", "overview.f.fallback", _label(labels, "overview.on" if fw["fallback"]["enabled"]
                                                         else "overview.off"))
        add("forwarding", "overview.f.fallbackUrl", fw["fallback"]["url"])
        add("forwarding", "overview.f.fallbackRequests", fw["fallback"]["requests"], fw["fallback"]["previous"])
        add("forwarding", "overview.f.rulesActive", fw["rules"]["active"])
        add("forwarding", "overview.f.rulesTotal", fw["rules"]["total"])
        add("forwarding", "overview.f.rulesRequests", fw["rules"]["requests"], fw["rules"]["previous"])
    gov = data.get("governance")
    if gov:
        for action, n in sorted(gov["changes"].items()):
            add("governance", "overview.g.change." + action, n)
        add("governance", "overview.g.signIns", gov["signIns"])
        add("governance", "overview.g.signInsRefused", gov["signInsRefused"])
        add("governance", "overview.g.exports", gov["exports"])
        for role, n in sorted(gov["users"].items()):
            add("governance", "overview.g.users." + role, n)
        add("governance", "overview.g.keysActive", gov["apiKeys"]["active"])
        add("governance", "overview.g.keysExpiring", gov["apiKeys"]["expiring"])
        add("governance", "overview.g.apiCalls", gov["apiCalls"]["calls"])
        add("governance", "overview.g.apiRefused", gov["apiCalls"]["refused"])
    return rows


def tables(data: dict, labels: dict) -> list[tuple[str, list[str], list[list]]]:
    """The overview's lists as (sheet name, header, rows)."""
    L = lambda code: _label(labels, code)          # noqa: E731
    out = [(L("overview.sheet.summary"), [L("overview.col.group"), L("overview.col.indicator"),
                                          L("overview.col.value"), L("overview.col.previous")],
            summary_rows(data, labels)),
           (L("overview.sheet.perDay"), [L("overview.col.day"), L("overview.scans"), L("overview.automated"),
                                         L("overview.notFound")],
            [[d["day"], d["scans"], d["automated"], d["notFound"]] for d in data["usage"]["perDay"]]),
           (L("overview.sheet.topCodes"), [L("overview.col.code"), L("overview.scans"), L("overview.notFound")],
            [[c["anchor"], c["scans"], c["notFound"]] for c in data["usage"]["topCodes"]]),
           (L("overview.sheet.unknown"), [L("overview.col.code"), L("overview.col.requests"),
                                          L("overview.col.keyRegistered")],
            [[c["anchor"], c["requests"], "" if c["registered"] is None else L("overview.yes" if c["registered"]
                                                                                 else "overview.no")]
             for c in data["quality"]["unknownCodes"]]),
           (L("overview.sheet.noLink"), [L("overview.col.code"), L("overview.col.requests")],
            [[c["anchor"], c["requests"]] for c in data["quality"]["noLinkCodes"]])]
    if data["records"].get("available"):
        out.append((L("overview.sheet.notScanned"), [L("overview.col.code")],
                    [[a] for a in data["records"]["notScannedSample"]]))
    if data.get("forwarding"):
        fw = data["forwarding"]
        out.append((L("overview.sheet.fallback"), [L("overview.col.code"), L("overview.col.requests")],
                    [[c["anchor"], c["requests"]] for c in fw["fallback"]["topCodes"]]))
        out.append((L("overview.sheet.byRule"), [L("overview.col.prefix"), L("overview.col.destination"),
                                                 L("overview.col.state"), L("overview.col.requests")],
                    [[r["prefix"], r["target"], L("overview.on" if r["enabled"] else "overview.off"), r["requests"]]
                     for r in fw["rules"]["byRule"]]))
    if data.get("governance"):
        out.append((L("overview.sheet.byUser"), [L("overview.col.user"), L("overview.col.changes")],
                    [list(p) for p in data["governance"]["byUser"]]))
    return out


def write_csv(data: dict, labels: dict) -> bytes:
    """The single figures, one per row (UTF-8 with BOM, ';', cells protected against formulas)."""
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    name, header, rows = tables(data, labels)[0]
    writer.writerow(header)
    for row in rows:
        writer.writerow([sheet.protect(c) for c in row])
    return ("﻿" + out.getvalue()).encode("utf-8")


def write_xlsx(data: dict, labels: dict) -> bytes:
    """One sheet per list: the single figures first, then per day, codes, users."""
    book = Workbook()
    book.remove(book.active)
    for name, header, rows in tables(data, labels):
        ws = book.create_sheet(sheet_title(name))
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([sheet.protect(c) if isinstance(c, str) else c for c in row])
        for i, _ in enumerate(header, start=1):
            width = max([len(str(header[i - 1]))] + [len(str(r[i - 1])) for r in rows if len(r) >= i] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def sheet_title(name: str) -> str:
    """A worksheet's name: at most 31 characters, none of []:*?/\\."""
    clean = "".join(" " if ch in "[]:*?/\\" else ch for ch in str(name)).strip()
    return (clean or "Sheet")[:31]
